"""
scripts/load_user_data.py — 基礎資料一鍵載入器（環工部主表 → B 棧 plant/item/spec/tag_mapping）

背景：舊資料搬遷卡在「撞唯一鍵/孤兒 FK」，決定基礎資料整個重建、走全新的路，不再遷就舊資料。
原本要環工部填四張表（廠區清單/項目清單/規格門檻/Tag 對應），評估後併成一張主表——廠區與
項目清單本來就能從規格門檻那張推導出來，填四張反而會發生「項目清單有、規格表漏了」對不起來
的狀況。本腳本讀那張主表（+ 選填的廠區設定表），推導並 upsert 出 plant/item/spec/tag_mapping。

主表欄位（表頭固定用中文，與環工部約定的格式，不可自行更改欄名）：
    廠區代號*／項目*／顯示名／單位／類型*／法規值／OOS／OOC／Alert／允收值／
    資料來源*／Tag名稱／排序
    （*=必填；一列＝一個廠區的一個監測項目）

選填第二張表「廠區設定」：欄位 廠區代號*／顯示排序／是否顯示(Y/N)。
  - .xlsx：多工作表時，主表用第一張或名為「主表」的那張，廠區設定用名為「廠區設定」的那張。
  - .csv：主表即 --file 指定的檔案；廠區設定表則讀同目錄下檔名為「廠區設定.csv」的檔案
    （若不存在則視為未提供，比照 xlsx 沒有那張工作表的行為：排序依主表首次出現順序、一律顯示）。

行為摘要（細節見各函式 docstring）：
  1. 推導主檔：從主表 distinct 出廠區與項目，自動建 plant/item。plant_id/item_id 非自增，
     已存在的沿用既有 id；新的從現有最大值 +1 開始遞增（plant_id 是 smallint，上限 32767，
     超過報錯中止）。
  2. 建 spec：門檻欄一律經 schemas.spec_schema.parse_spec_bound() 驗證後拆成 low/high/status
     （可解析→valid；空→na；'建置中'→building；pH 類項目必須雙邊，由該函式負責擋）。
  3. 建 tag_mapping：資料來源=SCADA 且 Tag名稱有值時才產生一列（target_field 固定 'value'，
     enabled 固定 True，source_table 由 --source-table 指定，預設 'kepware_sim'）。
  4. 冪等：以業務鍵（plant_no / item / (plant_no,item) / (source_table,tagname)）查找後
     upsert，重跑同一份檔案不會產生重複列或報錯。
  5. --dry-run：跑完整驗證＋（在記憶體中）試算異動，最後 rollback，完全不寫 DB。
  6. 逐列錯誤報告：任一列有問題時印出「第 N 列 [欄位]：原因」，不中斷整批——全部檢查完
     一次報告，最後彙總「成功 X 列、失敗 Y 列」。同一列若有多個欄位出錯，全部列出，
     該列所有異動（含連動的 tag_mapping）都不會寫入。
  6-1. 逐列**警告**（「⚠ 警告」前綴，獨立一段輸出，--dry-run 也照印）：門檻階梯不完整
     （填了 OOC 卻沒填 OOS、填了 Alert 卻沒填 OOC）時提醒該門檻不會生效。警告**不是錯誤**、
     不影響載入也不影響離開碼——「建置中」項目整排留空是合法的，只有「填一半」才可疑。
     詳見 check_threshold_ladder()。
  7. 支援 .csv（UTF-8，含 BOM 也讀得動）與 .xlsx（openpyxl）。

用法：
  python scripts/load_user_data.py --file 基礎資料主表.xlsx
  python scripts/load_user_data.py --file 基礎資料主表.csv --dry-run
  python scripts/load_user_data.py --file 基礎資料主表.csv --source-table kepware_sim

★ 「類型」欄（水質/空汙/雨水溝）驗證後寫入 item.category（2026-07-31 新增的欄位）。
  背景：現有程式碼判斷項目類別是用名稱字串比對（services/dispatch_service.py 的
  "VOC" in item、"雨水溝" in item，services/flow_service.py 的簽核 rtype 也是），
  新廠若有不叫 VOC 的空汙項目會被誤判成水質，連帶找錯簽核人、派報分錯類。
  本載入器把使用者填的正確類型存進 item.category，資料先就位；
  ★ 各處判斷邏輯改讀 category 取代字串比對，是後續待辦（見 docs/標準化與待調整清單.md）。
"""

import argparse
import csv
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple

# 讓腳本可以在 repo 根目錄外執行時仍然找得到 database_b / models_b（專案根目錄加入 sys.path）
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from database_b import BSessionLocal, create_all_b  # noqa: E402
from models_b import Item, Plant, Source, Spec, TagMapping  # noqa: E402
from schemas.spec_schema import (  # noqa: E402
    check_ooc_alert_hierarchy, parse_spec_bound,
)
from sqlalchemy import func, select  # noqa: E402

# ── 主表欄名常數（與環工部約定的格式，不可自行更改）────────────────────────
COL_PLANT_NO = "廠區代號"
COL_ITEM = "項目"
COL_DISPLAY = "顯示名"
COL_UNIT = "單位"
COL_TYPE = "類型"
COL_LAW = "法規值"
COL_OOS = "OOS"
COL_OOC = "OOC"
COL_ALERT = "Alert"
COL_RECV = "允收值"
COL_SOURCE = "資料來源"
COL_TAG = "Tag名稱"
COL_SEQ = "排序"

MAIN_COLUMNS = [
    COL_PLANT_NO, COL_ITEM, COL_DISPLAY, COL_UNIT, COL_TYPE, COL_LAW,
    COL_OOS, COL_OOC, COL_ALERT, COL_RECV, COL_SOURCE, COL_TAG, COL_SEQ,
]
REQUIRED_MAIN_COLUMNS = [COL_PLANT_NO, COL_ITEM, COL_TYPE, COL_SOURCE]

# 廠區設定表欄名
PLANT_CONF_PLANT_NO = "廠區代號"
PLANT_CONF_SORT = "顯示排序"
PLANT_CONF_SHOW = "是否顯示"

ITEM_TYPES = ("水質", "空汙", "雨水溝")
SOURCE_MAP = {"SCADA": 1, "CWMS": 2, "QA": 3}
SCADA_SOURCE_ID = SOURCE_MAP["SCADA"]

PLANT_ID_MAX = 32767  # SmallInteger 上限（models_b.Plant.plant_id）

_PH_RE = re.compile(r"^ph\d*$", re.IGNORECASE)
_BUILDING_TEXT = "建置中"


# ══════════════════════════════════════════════════════════════════════════
# 資料結構
# ══════════════════════════════════════════════════════════════════════════

@dataclass
class RowError:
    row: int
    field: str
    message: str

    def __str__(self) -> str:
        return f"第 {self.row} 列 [{self.field}]：{self.message}"


@dataclass
class RowWarning:
    """逐列警告：格式與 RowError 平行，但**不會**讓該列失敗、也不影響離開碼。

    輸出刻意加「⚠ 警告」前綴並在 CLI 獨立成一段，避免與「第 N 列 [欄位]：原因」的錯誤混淆。
    """
    row: int
    plant_no: Optional[str]
    item: Optional[str]
    field: str
    message: str

    def __str__(self) -> str:
        where = f"{self.plant_no or '?'}／{self.item or '?'}"
        return f"⚠ 警告 第 {self.row} 列 [{where}]（{self.field}）：{self.message}"


@dataclass
class ParsedRow:
    row: int
    plant_no: str
    item: str
    display_name: Optional[str]
    unit: Optional[str]
    law_text: Optional[str]
    oos: Tuple[Optional[Decimal], Optional[Decimal], str]
    ooc: Tuple[Optional[Decimal], Optional[Decimal], str]
    alert: Tuple[Optional[Decimal], Optional[Decimal], str]
    recv: Tuple[Optional[Decimal], Optional[Decimal], str]
    source_id: int
    tagname: Optional[str]
    seqno: Optional[int]
    category: Optional[str]   # 類型（水質/空汙/雨水溝）→ item.category


class _IdAllocator:
    """依現有 DB 最大值往上配號，同一次執行程序內連續遞增（避免同批次配到同一個新號）。"""

    def __init__(self, start: int, max_value: Optional[int] = None, label: str = "id"):
        self._next = start
        self._max_value = max_value
        self._label = label

    def take(self) -> int:
        val = self._next
        if self._max_value is not None and val > self._max_value:
            raise ValueError(
                f"{self._label} 配號超過上限 {self._max_value}（配到 {val}），"
                "請人工介入處理（例如先清理不再使用的舊代碼）！"
            )
        self._next += 1
        return val


# ══════════════════════════════════════════════════════════════════════════
# 純函式：欄位值處理
# ══════════════════════════════════════════════════════════════════════════

def _s(v: Any) -> Optional[str]:
    """儲存格值轉字串並 strip；None/空字串一律回 None（CSV 讀出必是 str，xlsx 可能是 int/float）。"""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def is_ph_item(item: str) -> bool:
    """項目名稱是否為 pH 類（雙邊規格）：'pH'/'pH1'/'pH2'... 皆算，不分大小寫。"""
    return bool(_PH_RE.match((item or "").strip()))


def _bound_triplet(raw: Any, field_name: str, is_ph: bool, item_name: str) -> Tuple[Optional[Decimal], Optional[Decimal], str]:
    """單一門檻欄位字串 → (low, high, status)。

    格式驗證一律重用 schemas.spec_schema.parse_spec_bound()（單邊/雙邊、pH 雙邊、最多兩位
    小數等規則都在裡面，不在這裡重寫）；'建置中' 是 parse_spec_bound 無法辨識的特例
    （它會判成格式錯誤丟例外），在這裡另外攔截、轉成 status='building'（比照
    migration/normalize/spec.py 的 _bounds_triplet 手法）。
    """
    s = _s(raw)
    try:
        vals = parse_spec_bound(s, field_name, is_ph, item_name)
    except ValueError:
        if s == _BUILDING_TEXT:
            return None, None, "building"
        raise

    if vals is None:
        return None, None, "na"
    if len(vals) == 1:
        return None, Decimal(str(vals[0])), "valid"
    return Decimal(str(vals[0])), Decimal(str(vals[1])), "valid"


def check_threshold_ladder(row_no: int, plant_no: Optional[str], item: Optional[str],
                            oos_status: str, ooc_status: str, alert_status: str) -> List[RowWarning]:
    """門檻階梯完整性檢查 → 回傳警告清單（**不是錯誤**，不會擋下該列）。

    背景（2026-07-31 主控查證）：`services/dashboard_service._calculate_light()` 的判定式是
        橙燈：`ooc is not None and oos is not None and ooc <= rvalue < oos`
        黃燈：`alert is not None and ooc is not None and alert < rvalue < ooc`
    也就是說**上一層門檻缺席時，這一層永遠不會亮**：
      - 填了 OOC 但 OOS 留空 → 橙燈永遠不會亮（區間的上界不存在）
      - 填了 Alert 但 OOC 留空 → 黃燈（Alert 那條）永遠不會亮
    使用者填了值卻完全不生效，是最難察覺的那種設定錯誤，因此在載入時就提醒。

    為什麼是警告不是錯誤：「尚未建置」的項目本來就整排門檻留空（合法，畫面顯示建置中），
    只有「填一半」才可疑；而且部分項目確實可能只想要某一層門檻。硬擋會讓合法資料載不進來，
    所以一律放行、只提醒，離開碼也不受影響。
    """
    warnings: List[RowWarning] = []
    if ooc_status == "valid" and oos_status != "valid":
        warnings.append(RowWarning(
            row_no, plant_no, item, f"{COL_OOC}/{COL_OOS}",
            f"已填 {COL_OOC} 但 {COL_OOS} 未填有效數值 → 這個 {COL_OOC} 永遠不會亮橙燈"
            f"（橙燈條件是「{COL_OOC} ≦ 讀值 < {COL_OOS}」，缺 {COL_OOS} 整條件就不成立）。"
            f"若此項目尚在建置中可忽略，否則請補上 {COL_OOS}。"))
    if alert_status == "valid" and ooc_status != "valid":
        warnings.append(RowWarning(
            row_no, plant_no, item, f"{COL_ALERT}/{COL_OOC}",
            f"已填 {COL_ALERT} 但 {COL_OOC} 未填有效數值 → 這個 {COL_ALERT} 永遠不會亮黃燈"
            f"（黃燈條件是「{COL_ALERT} < 讀值 < {COL_OOC}」，缺 {COL_OOC} 整條件就不成立）。"
            f"若此項目尚在建置中可忽略，否則請補上 {COL_OOC}。"))
    return warnings


def _parse_int(raw: Optional[str]) -> int:
    """'5'/'5.0' 皆可轉整數（xlsx 數字儲存格常帶浮點）；其餘丟 ValueError。"""
    return int(float(raw))


# ══════════════════════════════════════════════════════════════════════════
# 檔案讀取（.csv / .xlsx）
# ══════════════════════════════════════════════════════════════════════════

def _read_csv(path: str) -> List[Dict[str, Any]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames:
            reader.fieldnames = [(h or "").strip() for h in reader.fieldnames]
        return [dict(row) for row in reader]


def _read_xlsx_sheet(wb, sheet_name: str) -> List[Dict[str, Any]]:
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []
    header = [(str(h).strip() if h is not None else "") for h in header]
    out: List[Dict[str, Any]] = []
    for r in rows_iter:
        if all(c is None for c in r):
            continue  # 跳過整列空白（Excel 常見尾端空列）
        out.append({header[i]: r[i] for i in range(len(header)) if i < len(r)})
    return out


def read_main_and_config(file_path: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """讀主表 + 選填的廠區設定表，回傳 (main_rows, plant_config_rows)。"""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".csv":
        main_rows = _read_csv(file_path)
        conf_path = os.path.join(os.path.dirname(os.path.abspath(file_path)), "廠區設定.csv")
        conf_rows = _read_csv(conf_path) if os.path.exists(conf_path) else []
        return main_rows, conf_rows

    if ext in (".xlsx", ".xlsm"):
        import openpyxl  # 延後匯入：只有讀 xlsx 才需要

        wb = openpyxl.load_workbook(file_path, data_only=True)
        main_sheet = "主表" if "主表" in wb.sheetnames else wb.sheetnames[0]
        main_rows = _read_xlsx_sheet(wb, main_sheet)
        conf_rows = _read_xlsx_sheet(wb, "廠區設定") if "廠區設定" in wb.sheetnames else []
        return main_rows, conf_rows

    raise ValueError(f"不支援的檔案格式：{ext}（僅支援 .csv / .xlsx）")


# ══════════════════════════════════════════════════════════════════════════
# 驗證 + 解析（純邏輯，不連 DB）
# ══════════════════════════════════════════════════════════════════════════

def validate_main_rows(main_rows: List[Dict[str, Any]]) -> Tuple[List[ParsedRow], List[RowError], List[RowWarning]]:
    """逐列驗證主表。任一列有錯就整列排除（其他列不受影響，全部驗完才回傳）。

    回傳 (通過的列, 錯誤清單, 警告清單)。警告只針對**通過驗證的列**產生
    （沒通過的列本來就不會寫入，再報警告只是噪音；使用者修好錯誤後重跑自然會看到）。

    表頭列＝Excel 第 1 列，故第一筆資料列的 row 編號＝2，與使用者在 Excel 看到的列號一致。
    """
    parsed: List[ParsedRow] = []
    errors: List[RowError] = []
    warnings: List[RowWarning] = []

    # 同一 item 在不同列的 顯示名/單位 必須一致（item 是全域主檔，非逐廠區各自一份）；
    # 以「該 item 第一次出現的非空值」為基準，之後出現不同的非空值即視為衝突。
    item_seen: Dict[str, Dict[str, Any]] = {}

    for idx, raw_row in enumerate(main_rows):
        row_no = idx + 2
        row_errors: List[RowError] = []

        plant_no = _s(raw_row.get(COL_PLANT_NO))
        if not plant_no:
            row_errors.append(RowError(row_no, COL_PLANT_NO, "不可空白"))

        item = _s(raw_row.get(COL_ITEM))
        if not item:
            row_errors.append(RowError(row_no, COL_ITEM, "不可空白"))

        item_type = _s(raw_row.get(COL_TYPE))
        if not item_type:
            row_errors.append(RowError(row_no, COL_TYPE, "不可空白"))
        elif item_type not in ITEM_TYPES:
            row_errors.append(RowError(
                row_no, COL_TYPE, f"必須為「{'/'.join(ITEM_TYPES)}」其一，實際為「{item_type}」"))

        source_text = _s(raw_row.get(COL_SOURCE))
        source_id: Optional[int] = None
        if not source_text:
            row_errors.append(RowError(row_no, COL_SOURCE, "不可空白"))
        elif source_text.upper() not in SOURCE_MAP:
            row_errors.append(RowError(
                row_no, COL_SOURCE, f"必須為「{'/'.join(SOURCE_MAP)}」其一，實際為「{source_text}」"))
        else:
            source_id = SOURCE_MAP[source_text.upper()]

        display_name = _s(raw_row.get(COL_DISPLAY))
        unit = _s(raw_row.get(COL_UNIT))
        law_text = _s(raw_row.get(COL_LAW))
        tagname = _s(raw_row.get(COL_TAG))

        seq_raw = _s(raw_row.get(COL_SEQ))
        seqno: Optional[int] = None
        if seq_raw is not None:
            try:
                seqno = _parse_int(seq_raw)
            except ValueError:
                row_errors.append(RowError(row_no, COL_SEQ, f"必須為整數，實際為「{seq_raw}」"))

        # 門檻四欄：item 為空時無法判斷是否為 pH，略過（已有「項目不可空白」的錯誤在報了，
        # 不需要再疊加令人困惑的門檻錯誤）。
        oos: Tuple[Optional[Decimal], Optional[Decimal], str] = (None, None, "na")
        ooc: Tuple[Optional[Decimal], Optional[Decimal], str] = (None, None, "na")
        alert: Tuple[Optional[Decimal], Optional[Decimal], str] = (None, None, "na")
        recv: Tuple[Optional[Decimal], Optional[Decimal], str] = (None, None, "na")
        if item:
            is_ph = is_ph_item(item)
            for col, field_label, setter in (
                (COL_OOS, "OOS", "oos"),
                (COL_OOC, "OOC", "ooc"),
                (COL_ALERT, "Alert", "alert"),
                (COL_RECV, "允收值", "recv"),
            ):
                try:
                    triplet = _bound_triplet(raw_row.get(col), field_label, is_ph, item)
                except ValueError as e:
                    row_errors.append(RowError(row_no, col, str(e)))
                    continue
                if setter == "oos":
                    oos = triplet
                elif setter == "ooc":
                    ooc = triplet
                elif setter == "alert":
                    alert = triplet
                else:
                    recv = triplet

            # OOC/Alert 層遞驗證（2026-07-31 主控補上）：Alert 範圍必須包在 OOC 內
            # ——Alert 最先觸發（範圍最窄）、OOC 次之、OOS 最嚴重，此方向已由使用者於
            # C6 確認正確（見 docs/標準化與待調整清單.md），故此處套用 spec_schema 的
            # check_ooc_alert_hierarchy()。需求文件已向環工部承諾「填錯會擋下並指出哪一列」，
            # 沒有這段就是承諾跳票。兩邊門檻皆成功解析（status=valid）時才檢查。
            if ooc[2] == "valid" and alert[2] == "valid":
                ooc_vals = [float(v) for v in (ooc[0], ooc[1]) if v is not None]
                alert_vals = [float(v) for v in (alert[0], alert[1]) if v is not None]
                try:
                    check_ooc_alert_hierarchy(ooc_vals, alert_vals, is_ph)
                except ValueError as e:
                    row_errors.append(RowError(row_no, f"{COL_OOC}/{COL_ALERT}", str(e)))

        # 同一 item 顯示名/單位一致性檢查
        if item:
            info = item_seen.setdefault(item, {
                "display_row": None, "display": None, "unit_row": None, "unit": None,
            })
            if display_name is not None:
                if info["display"] is None:
                    info["display"] = display_name
                    info["display_row"] = row_no
                elif info["display"] != display_name:
                    row_errors.append(RowError(
                        row_no, COL_DISPLAY,
                        f"項目「{item}」顯示名與第 {info['display_row']} 列不一致"
                        f"（{info['display']} vs {display_name}），同一項目的顯示名須全檔一致"))
            if unit is not None:
                if info["unit"] is None:
                    info["unit"] = unit
                    info["unit_row"] = row_no
                elif info["unit"] != unit:
                    row_errors.append(RowError(
                        row_no, COL_UNIT,
                        f"項目「{item}」單位與第 {info['unit_row']} 列不一致"
                        f"（{info['unit']} vs {unit}），同一項目的單位須全檔一致"))

        if row_errors:
            errors.extend(row_errors)
            continue

        # 門檻階梯完整性提醒（警告，不擋列）：填了 OOC 沒填 OOS、填了 Alert 沒填 OOC
        warnings.extend(check_threshold_ladder(
            row_no, plant_no, item, oos[2], ooc[2], alert[2]))

        parsed.append(ParsedRow(
            row=row_no, plant_no=plant_no, item=item, display_name=display_name, unit=unit,
            law_text=law_text, oos=oos, ooc=ooc, alert=alert, recv=recv,
            source_id=source_id, tagname=tagname, seqno=seqno, category=item_type,
        ))

    return parsed, errors, warnings


def validate_plant_config(conf_rows: List[Dict[str, Any]]) -> Tuple[Dict[str, Dict[str, Any]], List[RowError]]:
    """廠區設定表驗證。回傳 {plant_no: {"sort": int|None, "is_show": bool}} + 錯誤清單。"""
    result: Dict[str, Dict[str, Any]] = {}
    errors: List[RowError] = []

    for idx, raw in enumerate(conf_rows):
        row_no = idx + 2
        plant_no = _s(raw.get(PLANT_CONF_PLANT_NO))
        if not plant_no:
            errors.append(RowError(row_no, PLANT_CONF_PLANT_NO, "不可空白"))
            continue

        sort: Optional[int] = None
        sort_raw = _s(raw.get(PLANT_CONF_SORT))
        if sort_raw is not None:
            try:
                sort = _parse_int(sort_raw)
            except ValueError:
                errors.append(RowError(row_no, PLANT_CONF_SORT, f"必須為整數，實際為「{sort_raw}」"))

        is_show = True
        show_raw = _s(raw.get(PLANT_CONF_SHOW))
        if show_raw is not None:
            if show_raw.upper() not in ("Y", "N"):
                errors.append(RowError(row_no, PLANT_CONF_SHOW, f"必須為 Y/N，實際為「{show_raw}」"))
            else:
                is_show = show_raw.upper() == "Y"

        result[plant_no] = {"sort": sort, "is_show": is_show}

    return result, errors


# ══════════════════════════════════════════════════════════════════════════
# 套用到 DB（upsert，冪等）
# ══════════════════════════════════════════════════════════════════════════

def apply_load(db, parsed_rows: List[ParsedRow], plant_config: Dict[str, Dict[str, Any]],
                source_table: str) -> Dict[str, int]:
    """把驗證通過的列 upsert 進 plant/item/spec/tag_mapping。回傳異動摘要 dict。

    冪等策略：一律先用業務鍵（plant_no / item / (plant_no,item) / (source_table,tagname)）
    查詢，找到就更新既有列（沿用既有 id），找不到才配新號新增——重跑同一份檔案不會產生
    重複列，也不會因為 UNIQUE 約束報錯。
    """
    now = datetime.now(timezone.utc)
    summary = {
        "plant_new": 0, "plant_reused": 0,
        "item_new": 0, "item_reused": 0,
        "spec_upsert": 0, "tag_upsert": 0,
    }

    # ── 資料來源目錄（系統常數，非使用者資料）──
    # spec.source_id 對 source 表有 FK。「基礎資料重建」的使用情境常常是對著一個乾淨的
    # 資料庫跑（正是本工具存在的理由），此時 source 表是空的，直接寫 spec 會撞 FK。
    # 這三筆是固定的系統常數（1=SCADA/2=CWMS/3=QA，與 scripts/seed_test_data._seed_catalog
    # 一致），由本載入器自行補齊，使用者不必為了跑這支而先去跑別的 seed 腳本。
    for sid, sname in ((1, "SCADA"), (2, "CWMS"), (3, "QA")):
        if db.execute(select(Source).where(Source.source_id == sid)).scalar_one_or_none() is None:
            db.add(Source(source_id=sid, name=sname))
    db.flush()

    # ── 廠區：依主表首次出現順序決定預設排序 ──
    plant_order: List[str] = []
    for row in parsed_rows:
        if row.plant_no not in plant_order:
            plant_order.append(row.plant_no)

    existing_plants = {p.plant_no: p for p in db.execute(select(Plant)).scalars().all()}
    plant_id_start = (db.execute(select(func.max(Plant.plant_id))).scalar() or 0) + 1
    plant_id_alloc = _IdAllocator(plant_id_start, PLANT_ID_MAX, "廠區代碼(plant_id)")

    plant_no_to_id: Dict[str, int] = {}
    for i, plant_no in enumerate(plant_order, start=1):
        conf = plant_config.get(plant_no, {})
        sort = conf.get("sort") if conf.get("sort") is not None else i
        is_show = conf.get("is_show", True)

        entry = existing_plants.get(plant_no)
        if entry is not None:
            entry.sort = sort
            entry.is_show = is_show
            plant_no_to_id[plant_no] = entry.plant_id
            summary["plant_reused"] += 1
        else:
            new_id = plant_id_alloc.take()
            db.add(Plant(plant_id=new_id, plant_no=plant_no, kind="normal", is_show=is_show, sort=sort))
            plant_no_to_id[plant_no] = new_id
            summary["plant_new"] += 1
    db.flush()

    # ── 項目：全域唯一，顯示名/單位取「第一次出現的非空值」（驗證階段已擋掉衝突）──
    item_order: List[str] = []
    item_meta: Dict[str, Tuple[Optional[str], Optional[str], Optional[str]]] = {}
    for row in parsed_rows:
        if row.item not in item_order:
            item_order.append(row.item)
            item_meta[row.item] = (None, None, None)
        dn, un, cat = item_meta[row.item]
        if dn is None and row.display_name is not None:
            dn = row.display_name
        if un is None and row.unit is not None:
            un = row.unit
        if cat is None and row.category:
            cat = row.category
        item_meta[row.item] = (dn, un, cat)

    existing_items = {i.item: i for i in db.execute(select(Item)).scalars().all()}
    item_id_start = (db.execute(select(func.max(Item.item_id))).scalar() or 0) + 1
    item_id_alloc = _IdAllocator(item_id_start, None, "項目代碼(item_id)")

    for item_name in item_order:
        display_name, unit, category = item_meta[item_name]
        entry = existing_items.get(item_name)
        if entry is not None:
            if display_name is not None:
                entry.display_name = display_name
            if unit is not None:
                entry.unit = unit
            if category is not None:
                entry.category = category
            summary["item_reused"] += 1
        else:
            new_id = item_id_alloc.take()
            db.add(Item(item_id=new_id, item=item_name, display_name=display_name,
                        unit=unit, category=category, is_active=True))
            summary["item_new"] += 1
    db.flush()

    # ── 規格：(plant_no, item) 複合主鍵 upsert；排序空白時依「廠區內出現順序」遞增配號 ──
    existing_specs = {(s.plant_no, s.item): s for s in db.execute(select(Spec)).scalars().all()}
    plant_seq_counter: Dict[str, int] = {}

    for row in parsed_rows:
        if row.seqno is not None:
            seqno = row.seqno
        else:
            plant_seq_counter[row.plant_no] = plant_seq_counter.get(row.plant_no, 0) + 1
            seqno = plant_seq_counter[row.plant_no]

        oos_low, oos_high, oos_status = row.oos
        ooc_low, ooc_high, ooc_status = row.ooc
        alert_low, alert_high, alert_status = row.alert
        recv_low, recv_high, recv_status = row.recv

        key = (row.plant_no, row.item)
        entry = existing_specs.get(key)
        if entry is None:
            entry = Spec(plant_no=row.plant_no, item=row.item)
            db.add(entry)
            existing_specs[key] = entry

        entry.law_text = row.law_text
        entry.oos_low, entry.oos_high, entry.oos_status = oos_low, oos_high, oos_status
        entry.ooc_low, entry.ooc_high, entry.ooc_status = ooc_low, ooc_high, ooc_status
        entry.alert_low, entry.alert_high, entry.alert_status = alert_low, alert_high, alert_status
        entry.recv_low, entry.recv_high, entry.recv_status = recv_low, recv_high, recv_status
        entry.source_id = row.source_id
        entry.seqno = seqno
        entry.updated_at = now
        summary["spec_upsert"] += 1
    db.flush()

    # ── Tag 對應：資料來源=SCADA 且 Tag名稱有值才產生，(source_table,tagname) upsert ──
    existing_tags = {
        (t.source_table, t.tagname): t
        for t in db.execute(select(TagMapping).where(TagMapping.source_table == source_table)).scalars().all()
    }

    for row in parsed_rows:
        if row.source_id != SCADA_SOURCE_ID or not row.tagname:
            continue
        key = (source_table, row.tagname)
        entry = existing_tags.get(key)
        if entry is None:
            entry = TagMapping(source_table=source_table, tagname=row.tagname)
            db.add(entry)
            existing_tags[key] = entry
        entry.plant_no = row.plant_no
        entry.item = row.item
        entry.target_field = "value"
        entry.enabled = True
        summary["tag_upsert"] += 1
    db.flush()

    return summary


# ══════════════════════════════════════════════════════════════════════════
# CLI
# ══════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(
        description="基礎資料一鍵載入器：讀環工部主表（CSV/Excel），重建 B 棧 plant/item/spec/tag_mapping。")
    parser.add_argument("--file", required=True, help="主表檔案路徑（.csv 或 .xlsx）")
    parser.add_argument("--dry-run", action="store_true", help="只驗證與印出異動摘要，完全不寫入資料庫")
    parser.add_argument("--source-table", default="kepware_sim",
                         help="tag_mapping.source_table（預設 kepware_sim，即模擬 A 端表）")
    args = parser.parse_args()

    if not os.path.exists(args.file):
        print(f"[錯誤] 找不到檔案：{args.file}")
        sys.exit(1)

    print(f"[load_user_data] 讀取檔案：{args.file}")
    try:
        main_rows, conf_rows = read_main_and_config(args.file)
    except Exception as e:
        print(f"[錯誤] 讀檔失敗：{e}")
        sys.exit(1)

    if not main_rows:
        print("[錯誤] 主表沒有任何資料列（或找不到主表工作表/表頭），請確認檔案內容。")
        sys.exit(1)

    missing_cols = [c for c in REQUIRED_MAIN_COLUMNS if c not in main_rows[0].keys()]
    if missing_cols:
        print(f"[錯誤] 主表缺少必要欄位：{'、'.join(missing_cols)}")
        print(f"       主表欄位應為：{'、'.join(MAIN_COLUMNS)}")
        sys.exit(1)

    print(f"[load_user_data] 主表共 {len(main_rows)} 列資料"
          + (f"、廠區設定表 {len(conf_rows)} 列" if conf_rows else "（未提供廠區設定表，將依主表首次出現順序排序、一律顯示）"))

    parsed, row_errors, row_warnings = validate_main_rows(main_rows)
    plant_config, conf_errors = validate_plant_config(conf_rows)
    all_errors = row_errors + conf_errors

    if all_errors:
        print(f"\n[驗證錯誤] 共 {len(all_errors)} 筆：")
        for err in all_errors:
            print(f"  {err}")

    # 警告獨立成一段（--dry-run 也照印）：不影響載入、不影響離開碼，但填了不生效的設定
    # 不該靜靜通過，見 check_threshold_ladder()。
    if row_warnings:
        print(f"\n[警告] 共 {len(row_warnings)} 筆（不影響載入，資料仍會寫入，但請確認是否為預期設定）：")
        for warn in row_warnings:
            print(f"  {warn}")

    ok_count = len(parsed)
    fail_rows = {e.row for e in row_errors}
    fail_count = len(fail_rows)
    print(f"\n[驗證結果] 成功 {ok_count} 列、失敗 {fail_count} 列、警告 {len(row_warnings)} 筆"
          f"（主表共 {len(main_rows)} 列）")

    if not parsed:
        print("[load_user_data] 沒有任何一列通過驗證，沒有可寫入的資料，結束。")
        sys.exit(1)

    create_all_b()
    db = BSessionLocal()
    try:
        summary = apply_load(db, parsed, plant_config, args.source_table)

        if args.dry_run:
            db.rollback()
            print("\n[DRY-RUN] 完全未寫入資料庫，以下為預計異動摘要：")
        else:
            db.commit()
            print("\n[load_user_data] 已寫入資料庫，異動摘要：")

        print(f"  廠區（plant）：新增 {summary['plant_new']} 筆、沿用既有 {summary['plant_reused']} 筆")
        print(f"  項目（item）：新增 {summary['item_new']} 筆、沿用既有 {summary['item_reused']} 筆")
        print(f"  規格（spec）：新增/更新 {summary['spec_upsert']} 筆")
        print(f"  Tag 對應（tag_mapping，source_table={args.source_table}）：新增/更新 {summary['tag_upsert']} 筆")
    except ValueError as e:
        db.rollback()
        print(f"\n[錯誤] {e}")
        sys.exit(1)
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

    sys.exit(0 if not all_errors else 2)


if __name__ == "__main__":
    main()
